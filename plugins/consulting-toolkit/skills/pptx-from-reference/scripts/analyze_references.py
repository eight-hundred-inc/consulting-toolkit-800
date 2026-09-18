#!/usr/bin/env python3
"""参照 pptx デッキを全件解析し、デザイン仕様（JSON）と要約（Markdown）を出力する。

使い方:
    python3 analyze_references.py [--decks DIR] [--outdir DIR] [--max-slides N] [--masters-only]

--decks を省略した場合、本スクリプトの親ディレクトリ配下の reference-decks/ を
再帰的に走査し、見つかった .pptx / .potx を **すべて** 解析対象にする。
出力は outdir/deck-spec.json（全量）と outdir/deck-spec.md（エージェント読解用の要約）。

--masters-only を付けると **スライドマスターのみ参照モード** になる:
既存スライドの中身は一切解析せず、スライドマスター・レイアウト・テーマ
（配色・フォント・マスター既定書式・マスター上のロゴ／フッタ等）だけを抽出する。
参照デッキの完成スライドを真似ず、マスターとテーマだけ継承したい場合に使う。
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:  # pragma: no cover
    sys.exit("python-pptx が必要です: pip install python-pptx")

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


# --------------------------------------------------------------------------- utils
def inch(emu) -> float | None:
    if emu is None:
        return None
    return round(Emu(int(emu)).inches, 3)


def pt(size) -> float | None:
    if size is None:
        return None
    try:
        return round(size.pt, 1)
    except AttributeError:
        return None


def hexval(color_el) -> str | None:
    """<a:srgbClr val="..."/> / <a:sysClr lastClr="..."/> / <a:schemeClr val="..."/> を読む。"""
    if color_el is None:
        return None
    srgb = color_el.find("a:srgbClr", NS)
    if srgb is not None:
        return srgb.get("val", "").upper() or None
    sysc = color_el.find("a:sysClr", NS)
    if sysc is not None:
        return (sysc.get("lastClr") or "").upper() or None
    scheme = color_el.find("a:schemeClr", NS)
    if scheme is not None:
        return "scheme:" + scheme.get("val", "")
    if color_el.tag.endswith("}srgbClr"):
        return color_el.get("val", "").upper() or None
    return None


def safe_rgb(color) -> str | None:
    try:
        if color is None or color.type is None:
            return None
        return str(color.rgb).upper()
    except Exception:
        try:
            return "scheme:%s" % color.theme_color
        except Exception:
            return None


# --------------------------------------------------------------------------- theme
def read_theme(prs) -> dict:
    """テーマ（配色・フォント）をパッケージから直接読む。"""
    out: dict = {"colors": {}, "fonts": {}}
    theme_part = None
    for part in prs.part.package.iter_parts():
        if "theme" in str(part.partname) and str(part.partname).endswith(".xml"):
            theme_part = part
            break
    if theme_part is None:
        return out
    try:
        from lxml import etree

        root = etree.fromstring(theme_part.blob)
    except Exception:
        return out

    scheme = root.find(".//a:clrScheme", NS)
    if scheme is not None:
        for child in scheme:
            name = etree.QName(child).localname
            out["colors"][name] = hexval(child)

    fonts = root.find(".//a:fontScheme", NS)
    if fonts is not None:
        for kind in ("majorFont", "minorFont"):
            el = fonts.find("a:%s" % kind, NS)
            if el is None:
                continue
            latin = el.find("a:latin", NS)
            ea = el.find("a:ea", NS)
            jpan = el.find('a:font[@script="Jpan"]', NS)
            out["fonts"][kind] = {
                "latin": latin.get("typeface") if latin is not None else None,
                "ea": (ea.get("typeface") if ea is not None else None)
                or (jpan.get("typeface") if jpan is not None else None),
            }
    return out


# --------------------------------------------------------------------------- shapes
# --------------------------------------------------------------------------- inheritance
def _defrpr_props(defrpr) -> dict:
    """<a:defRPr> から sz/b/color/latin/ea を取り出す。"""
    if defrpr is None:
        return {}
    out = {}
    sz = defrpr.get("sz")
    if sz:
        out["size_pt"] = round(int(sz) / 100.0, 1)
    if defrpr.get("b") is not None:
        out["bold"] = defrpr.get("b") in ("1", "true")
    fill = defrpr.find("a:solidFill", NS)
    if fill is not None:
        out["color"] = hexval(fill)
    latin = defrpr.find("a:latin", NS)
    if latin is not None:
        out["latin_font"] = latin.get("typeface")
    ea = defrpr.find("a:ea", NS)
    if ea is not None:
        out["ea_font"] = ea.get("typeface")
    return out


_PH_STYLE_MAP = {
    "TITLE": "p:titleStyle", "CENTER_TITLE": "p:titleStyle",
    "SUBTITLE": "p:bodyStyle", "BODY": "p:bodyStyle", "OBJECT": "p:bodyStyle",
}


def inherited_defaults(shape, slide, level: int = 0) -> dict:
    """プレースホルダが継承する既定書式（レイアウト→マスター）を解決する。"""
    if not shape.is_placeholder:
        return {}
    try:
        fmt = shape.placeholder_format
        idx, ph_type = fmt.idx, str(fmt.type).split(" ")[0]
    except Exception:
        return {}

    props: dict = {}
    layout = getattr(slide, "slide_layout", None)
    for source in (layout, getattr(layout, "slide_master", None)):
        if source is None:
            continue
        for ph in getattr(source, "placeholders", []):
            try:
                if ph.placeholder_format.idx != idx:
                    continue
            except Exception:
                continue
            lst = ph._element.find(".//a:lstStyle/a:lvl%dpPr/a:defRPr" % (level + 1), NS)
            for k, v in _defrpr_props(lst).items():
                props.setdefault(k, v)

    master = getattr(layout, "slide_master", None)
    if master is not None:
        style_tag = _PH_STYLE_MAP.get(ph_type, "p:otherStyle")
        node = master._element.find(".//p:txStyles/%s/a:lvl%dpPr/a:defRPr" % (style_tag, level + 1), NS)
        for k, v in _defrpr_props(node).items():
            props.setdefault(k, v)
    return props


def read_runs(shape, slide=None) -> list[dict]:
    if not getattr(shape, "has_text_frame", False):
        return []
    runs = []
    inherit_cache: dict = {}
    for p_i, para in enumerate(shape.text_frame.paragraphs):
        for run in para.runs:
            text = run.text
            if not text.strip():
                continue
            rpr = run._r.find("a:rPr", NS)
            ea_font = None
            if rpr is not None:
                ea = rpr.find("a:ea", NS)
                ea_font = ea.get("typeface") if ea is not None else None
            if slide is not None and shape.is_placeholder:
                if para.level not in inherit_cache:
                    inherit_cache[para.level] = inherited_defaults(shape, slide, para.level)
                inh = inherit_cache[para.level]
            else:
                inh = {}
            runs.append(
                {
                    "para": p_i,
                    "text": text,
                    "size_pt": pt(run.font.size) or inh.get("size_pt"),
                    "bold": run.font.bold or inh.get("bold"),
                    "italic": run.font.italic,
                    "latin_font": run.font.name or inh.get("latin_font"),
                    "ea_font": ea_font or inh.get("ea_font"),
                    "color": safe_rgb(run.font.color) or inh.get("color"),
                    "align": str(para.alignment) if para.alignment is not None else None,
                    "level": para.level,
                }
            )
    return runs


def read_fill(shape) -> str | None:
    try:
        fill = shape.fill
        if fill.type is None:
            return None
        ftype = str(fill.type)
        if "SOLID" in ftype or fill.type == 1:
            return safe_rgb(fill.fore_color)
        if "BACKGROUND" in ftype or fill.type == 5:
            return "none"
        return ftype
    except Exception:
        return None


def read_line(shape) -> dict | None:
    try:
        line = shape.line
        color = safe_rgb(line.color)
        width = pt(line.width)
        if color is None and width is None:
            return None
        return {"color": color, "width_pt": width}
    except Exception:
        return None


def read_table(shape) -> dict | None:
    if not getattr(shape, "has_table", False):
        return None
    tbl = shape.table
    rows = []
    for r in tbl.rows:
        rows.append([c.text.replace("\n", " / ")[:60] for c in r.cells])
    return {
        "rows": len(tbl.rows),
        "cols": len(tbl.columns),
        "col_widths_in": [inch(c.width) for c in tbl.columns],
        "first_row_header": tbl.first_row,
        "preview": rows[:4],
    }


def read_shape(shape, slide=None, depth: int = 0) -> dict:
    info = {
        "name": shape.name,
        "type": str(shape.shape_type),
        "x_in": inch(shape.left),
        "y_in": inch(shape.top),
        "w_in": inch(shape.width),
        "h_in": inch(shape.height),
        "rotation": round(shape.rotation, 1) if shape.rotation else None,
    }
    if shape.is_placeholder:
        try:
            info["placeholder"] = {
                "idx": shape.placeholder_format.idx,
                "type": str(shape.placeholder_format.type),
            }
        except Exception:
            pass

    try:
        auto = shape._element.find(".//a:prstGeom", NS)
        if auto is not None:
            info["geom"] = auto.get("prst")
    except Exception:
        pass

    fill = read_fill(shape)
    if fill:
        info["fill"] = fill
    line = read_line(shape)
    if line:
        info["line"] = line

    runs = read_runs(shape, slide)
    if runs:
        info["runs"] = runs
        info["text"] = " ".join(r["text"] for r in runs)[:300]

    table = read_table(shape)
    if table:
        info["table"] = table

    if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
        info["picture"] = True

    if getattr(shape, "has_chart", False):
        try:
            info["chart"] = str(shape.chart.chart_type)
        except Exception:
            info["chart"] = True

    if shape.shape_type is not None and "GROUP" in str(shape.shape_type) and depth < 3:
        info["children"] = [read_shape(s, slide, depth + 1) for s in shape.shapes]
    return info


# --------------------------------------------------------------------------- masters
def read_master_txstyles(master) -> dict:
    """マスターの txStyles（タイトル／本文の既定書式）を lvl1 について読む。"""
    out = {}
    for label, tag in (("titleStyle", "p:titleStyle"), ("bodyStyle", "p:bodyStyle"),
                       ("otherStyle", "p:otherStyle")):
        node = master._element.find(".//p:txStyles/%s/a:lvl1pPr/a:defRPr" % tag, NS)
        props = _defrpr_props(node)
        if props:
            out[label] = props
    return out


def analyze_masters(prs) -> list[dict]:
    """スライドマスター自体を実測する（マスター上のロゴ・フッタ・帯・既定書式）。"""
    masters = []
    for m_i, master in enumerate(prs.slide_masters, start=1):
        bg = None
        try:
            bg_el = master._element.find(".//p:bg", NS)
            if bg_el is not None:
                bg = hexval(bg_el.find(".//a:solidFill", NS))
        except Exception:
            pass
        masters.append(
            {
                "no": m_i,
                "name": master.name or ("master %d" % m_i),
                "background": bg,
                "tx_styles": read_master_txstyles(master),
                "shapes": [read_shape(s) for s in master.shapes],
            }
        )
    return masters


# --------------------------------------------------------------------------- deck
def analyze_deck(path: Path, max_slides: int | None, masters_only: bool = False) -> dict:
    prs = Presentation(str(path))
    deck = {
        "file": path.name,
        "path": str(path),
        "mode": "masters-only" if masters_only else "full",
        "slide_w_in": inch(prs.slide_width),
        "slide_h_in": inch(prs.slide_height),
        "slide_count": len(prs.slides),
        "theme": read_theme(prs),
        "masters": analyze_masters(prs),
        "layouts": [],
        "slides": [],
    }
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            entry = {
                "name": layout.name,
                "placeholders": [
                    {
                        "idx": ph.placeholder_format.idx,
                        "type": str(ph.placeholder_format.type),
                        "x_in": inch(ph.left),
                        "y_in": inch(ph.top),
                        "w_in": inch(ph.width),
                        "h_in": inch(ph.height),
                    }
                    for ph in layout.placeholders
                ],
                "shape_count": len(layout.shapes),
            }
            if masters_only:
                # マスターのみ参照モードではレイアウト上の装飾（帯・線・ロゴ等）も
                # 座標根拠になるため全図形をダンプする
                entry["shapes"] = [read_shape(s) for s in layout.shapes]
            deck["layouts"].append(entry)

    if masters_only:
        return deck

    slides = list(prs.slides)
    if max_slides:
        slides = slides[:max_slides]
    for i, slide in enumerate(slides, start=1):
        bg = None
        try:
            bg_el = slide._element.find(".//p:bg", NS)
            if bg_el is not None:
                bg = hexval(bg_el.find(".//a:solidFill", NS))
        except Exception:
            pass
        deck["slides"].append(
            {
                "no": i,
                "layout": slide.slide_layout.name,
                "background": bg,
                "shapes": [read_shape(s, slide) for s in slide.shapes],
            }
        )
    return deck


# --------------------------------------------------------------------------- stats
def iter_tree(shapes):
    stack = list(shapes)
    while stack:
        sh = stack.pop()
        yield sh
        stack.extend(sh.get("children", []))


def walk_shapes(deck):
    """統計対象の図形を列挙する。full はスライド、masters-only はマスター＋レイアウト。"""
    if deck["slides"]:
        for sl in deck["slides"]:
            for sh in iter_tree(sl["shapes"]):
                yield sl, sh
    else:
        for m in deck.get("masters", []):
            for sh in iter_tree(m["shapes"]):
                yield m, sh
        for lay in deck.get("layouts", []):
            for sh in iter_tree(lay.get("shapes", [])):
                yield lay, sh


def deck_stats(deck: dict) -> dict:
    colors, fonts, sizes, geoms = Counter(), Counter(), Counter(), Counter()
    layouts = Counter(sl["layout"] for sl in deck["slides"])
    for sl, sh in walk_shapes(deck):
        if sh.get("fill") and not str(sh["fill"]).startswith(("scheme", "none")):
            colors["fill:%s" % sh["fill"]] += 1
        if sh.get("geom"):
            geoms[sh["geom"]] += 1
        for run in sh.get("runs", []):
            if run["color"] and not str(run["color"]).startswith("scheme"):
                colors["text:%s" % run["color"]] += 1
            for key in ("latin_font", "ea_font"):
                if run[key]:
                    fonts["%s(%s)" % (run[key], key[:5])] += 1
            if run["size_pt"]:
                sizes[run["size_pt"]] += 1
    return {
        "colors": colors.most_common(24),
        "fonts": fonts.most_common(12),
        "font_sizes": sorted(sizes.items(), key=lambda kv: -kv[1])[:16],
        "geoms": geoms.most_common(12),
        "layout_usage": layouts.most_common(),
    }


def recurring_shapes(deck: dict) -> list[dict]:
    """全スライドで同じ位置に繰り返し現れる要素（ヘッダ・フッタ・ページ番号・ロゴ）を検出。"""
    if not deck["slides"]:  # masters-only: 反復要素はマスター実測ダンプ側で示す
        return []
    buckets = defaultdict(list)
    for sl, sh in walk_shapes(deck):
        if sh["x_in"] is None:
            continue
        key = (round(sh["x_in"], 1), round(sh["y_in"], 1), round(sh["w_in"] or 0, 1), sh["type"])
        buckets[key].append((sl["no"], sh))
    n = max(len(deck["slides"]), 1)
    out = []
    for key, items in buckets.items():
        if len(items) < max(3, n * 0.4):
            continue
        sample = items[0][1]
        out.append(
            {
                "occurrences": len(items),
                "coverage_pct": round(100 * len(items) / n),
                "x_in": key[0], "y_in": key[1], "w_in": key[2], "h_in": sample.get("h_in"),
                "type": key[3],
                "fill": sample.get("fill"),
                "sample_text": (sample.get("text") or "")[:60],
                "sample_size_pt": sample["runs"][0]["size_pt"] if sample.get("runs") else None,
            }
        )
    return sorted(out, key=lambda d: -d["occurrences"])[:20]


# --------------------------------------------------------------------------- markdown
def fmt_shape(sh: dict, indent: str = "") -> list[str]:
    bits = []
    box = "x={x} y={y} w={w} h={h}".format(
        x=sh["x_in"], y=sh["y_in"], w=sh["w_in"], h=sh["h_in"]
    )
    head = "%s- `%s` %s %s" % (indent, sh["type"], box, "[%s]" % sh["geom"] if sh.get("geom") else "")
    if sh.get("placeholder"):
        head += " ph=%s/%s" % (sh["placeholder"]["idx"], sh["placeholder"]["type"])
    if sh.get("fill"):
        head += " fill=%s" % sh["fill"]
    if sh.get("line"):
        head += " line=%s/%spt" % (sh["line"].get("color"), sh["line"].get("width_pt"))
    bits.append(head)
    for run in sh.get("runs", [])[:8]:
        bits.append(
            '%s    - %spt%s %s %s/%s : "%s"'
            % (
                indent,
                run["size_pt"],
                " bold" if run["bold"] else "",
                run["color"] or "-",
                run["latin_font"] or "-",
                run["ea_font"] or "-",
                run["text"].replace("\n", " ")[:90],
            )
        )
    if sh.get("table"):
        t = sh["table"]
        bits.append("%s    - table %dx%d widths=%s" % (indent, t["rows"], t["cols"], t["col_widths_in"]))
        for row in t["preview"]:
            bits.append("%s      | %s |" % (indent, " | ".join(row)))
    for child in sh.get("children", [])[:12]:
        bits.extend(fmt_shape(child, indent + "  "))
    return bits


def deck_markdown(deck: dict, stats: dict, recur: list[dict], slide_detail: int) -> str:
    masters_only = deck.get("mode") == "masters-only"
    L: list[str] = []
    L.append("## %s" % deck["file"])
    L.append("")
    L.append(
        "- キャンバス: %s × %s inch / スライド数: %d%s"
        % (deck["slide_w_in"], deck["slide_h_in"], deck["slide_count"],
           "（masters-only: 既存スライドは解析対象外）" if masters_only else "")
    )
    th = deck["theme"]
    if th.get("fonts"):
        for kind, f in th["fonts"].items():
            L.append("- テーマフォント %s: latin=`%s` / ea=`%s`" % (kind, f.get("latin"), f.get("ea")))
    if th.get("colors"):
        L.append("- テーマ配色: " + ", ".join("%s=%s" % (k, v) for k, v in th["colors"].items()))
    L.append("")
    L.append("### 実使用カラー（出現頻度順%s）" % ("・マスター/レイアウト由来" if masters_only else ""))
    L.append("")
    L.append("| 色 | 出現数 |")
    L.append("|---|---|")
    for name, cnt in stats["colors"]:
        L.append("| %s | %d |" % (name, cnt))
    L.append("")
    L.append("### 実使用フォント / サイズ")
    L.append("")
    L.append("- フォント: " + ", ".join("%s×%d" % (k, v) for k, v in stats["fonts"]))
    L.append("- サイズ(pt): " + ", ".join("%s×%d" % (k, v) for k, v in stats["font_sizes"]))
    L.append("- 図形geom: " + ", ".join("%s×%d" % (k, v) for k, v in stats["geoms"]))
    L.append("")
    L.append("### スライドレイアウト（マスター）")
    L.append("")
    L.append("| レイアウト名 | プレースホルダ | 使用スライド数 |")
    L.append("|---|---|---|")
    usage = dict(stats["layout_usage"])
    for lay in deck["layouts"]:
        phs = ", ".join("%s@%s,%s" % (p["type"].split(" ")[0], p["x_in"], p["y_in"]) for p in lay["placeholders"])
        L.append("| %s | %s | %s |" % (lay["name"], phs or "-",
                                       "-" if masters_only else usage.get(lay["name"], 0)))
    L.append("")
    if masters_only:
        for m in deck.get("masters", []):
            if m.get("tx_styles"):
                L.append("### マスター既定書式（txStyles lvl1）— %s" % m["name"])
                L.append("")
                for label, props in m["tx_styles"].items():
                    L.append("- %s: %s" % (label, ", ".join("%s=%s" % kv for kv in props.items())))
                L.append("")
        L.append("### スライドマスター 実測ダンプ")
        L.append("")
        for m in deck.get("masters", []):
            L.append("#### master %d — %s%s"
                     % (m["no"], m["name"], " / bg=%s" % m["background"] if m.get("background") else ""))
            for sh in m["shapes"]:
                L.extend(fmt_shape(sh))
            L.append("")
        L.append("### レイアウト別 実測ダンプ")
        L.append("")
        for lay in deck["layouts"]:
            if not lay.get("shapes"):
                continue
            L.append("#### layout — %s" % lay["name"])
            for sh in lay["shapes"]:
                L.extend(fmt_shape(sh))
            L.append("")
        return "\n".join(L)
    if recur:
        L.append("### 反復要素（ヘッダ / フッタ / ページ番号 / ロゴ の候補）")
        L.append("")
        L.append("| 出現率 | 種別 | x | y | w | h | fill | サンプル文字 | pt |")
        L.append("|---|---|---|---|---|---|---|---|---|")
        for r in recur:
            L.append(
                "| %d%% | %s | %s | %s | %s | %s | %s | %s | %s |"
                % (r["coverage_pct"], r["type"], r["x_in"], r["y_in"], r["w_in"], r["h_in"],
                   r["fill"], r["sample_text"].replace("|", "/"), r["sample_size_pt"])
            )
        L.append("")
    L.append("### スライド別 実測ダンプ（先頭 %d 枚）" % slide_detail)
    L.append("")
    for sl in deck["slides"][:slide_detail]:
        L.append("#### slide %d — layout: %s%s" % (sl["no"], sl["layout"], " / bg=%s" % sl["background"] if sl["background"] else ""))
        for sh in sl["shapes"]:
            L.extend(fmt_shape(sh))
        L.append("")
    return "\n".join(L)


# --------------------------------------------------------------------------- main
def main() -> int:
    here = Path(__file__).resolve().parent
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--decks", default=str(here.parent / "reference-decks"),
                    help="参照 pptx を置いたディレクトリ（既定: スキル配下の reference-decks/）")
    ap.add_argument("--outdir", default=".", help="出力先ディレクトリ")
    ap.add_argument("--max-slides", type=int, default=0, help="1デッキあたりの解析スライド上限（0=全件）")
    ap.add_argument("--slide-detail", type=int, default=12, help="Markdown に実測ダンプするスライド枚数")
    ap.add_argument("--masters-only", action="store_true",
                    help="スライドマスターのみ参照モード: 既存スライドは解析せず、"
                         "マスター・レイアウト・テーマだけを抽出する")
    args = ap.parse_args()

    decks_dir = Path(args.decks).expanduser().resolve()
    if not decks_dir.is_dir():
        print("参照ディレクトリがありません: %s" % decks_dir, file=sys.stderr)
        return 2

    files = sorted(
        p for p in decks_dir.rglob("*")
        if p.suffix.lower() in (".pptx", ".potx") and not p.name.startswith("~$")
    )
    if not files:
        print("参照 pptx が 1 件もありません。%s に pptx を配備してください。" % decks_dir, file=sys.stderr)
        return 3

    outdir = Path(args.outdir).expanduser().resolve()
    outdir.mkdir(parents=True, exist_ok=True)

    all_decks, md_parts = [], []
    md_parts.append("# 参照デッキ デザイン仕様（自動抽出%s）" % ("・マスターのみ参照モード" if args.masters_only else ""))
    md_parts.append("")
    md_parts.append("参照元: `%s` / 対象 %d ファイル" % (decks_dir, len(files)))
    if args.masters_only:
        md_parts.append("")
        md_parts.append("> masters-only: 既存スライドの中身は解析していない。"
                        "座標・配色の根拠はスライドマスター・レイアウト・テーマのみ。")
    md_parts.append("")
    md_parts.append("| # | ファイル | スライド数 | キャンバス |")
    md_parts.append("|---|---|---|---|")

    parsed = []
    for i, f in enumerate(files, start=1):
        try:
            deck = analyze_deck(f, args.max_slides or None, masters_only=args.masters_only)
        except Exception as e:  # 壊れたファイルがあっても残りを処理する
            print("!! 解析失敗 %s: %s" % (f.name, e), file=sys.stderr)
            continue
        stats = deck_stats(deck)
        recur = recurring_shapes(deck)
        deck["stats"] = stats
        deck["recurring_shapes"] = recur
        parsed.append((deck, stats, recur))
        all_decks.append(deck)
        md_parts.append("| %d | %s | %d | %s×%s |" % (i, f.name, deck["slide_count"], deck["slide_w_in"], deck["slide_h_in"]))

    md_parts.append("")
    for deck, stats, recur in parsed:
        md_parts.append(deck_markdown(deck, stats, recur, args.slide_detail))
        md_parts.append("")

    json_path = outdir / "deck-spec.json"
    md_path = outdir / "deck-spec.md"
    json_path.write_text(json.dumps({"decks": all_decks}, ensure_ascii=False, indent=1), encoding="utf-8")
    md_path.write_text("\n".join(md_parts), encoding="utf-8")
    print("OK: %d デッキを解析 -> %s / %s" % (len(all_decks), md_path, json_path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
