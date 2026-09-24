#!/usr/bin/env python3
"""参照デッキをテンプレートに使って pptx を直接組み立てるヘルパー。

設計方針:
  参照 pptx を **そのままテンプレートとして開き**、既存スライドだけを削除する。
  これによりテーマ配色・フォントスキーム・スライドマスター・レイアウト・
  スライドサイズが 1 バイトも変わらずに引き継がれる。以降は本モジュールの
  ヘルパーで、参照デッキから採寸した座標どおりに図形を置いていく。

    from deck_kit import *
    prs = open_template("reference-decks/foo.pptx")
    st  = load_style("deck-spec.json")          # 任意: 抽出済みの配色/フォント
    s   = add_slide(prs, "1_本文")
    set_placeholder(s, "TITLE", "見出し")
    add_text(s, "本文", 0.37, 1.2, 12.6, 0.4, size=14)
    save(prs, "output.pptx")
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

__all__ = [
    "open_template", "new_from_scratch", "save", "load_style", "Style",
    "list_layouts", "layout_by_name", "add_slide", "set_placeholder",
    "drop_empty_placeholders", "set_slide_background",
    "add_text", "add_bullets", "add_rect", "add_line", "add_table", "add_picture",
    "set_run_fonts", "no_shadow", "rgb", "fit_size", "text_width_em",
    "ALIGN", "ANCHOR", "Inches", "Pt", "MSO_SHAPE",
]

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
_PLAIN_TABLE_STYLE = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid


# --------------------------------------------------------------------------- style
class Style(dict):
    """deck-spec.json から起こした配色・フォント・キャンバス情報。属性でも引ける。"""

    def __getattr__(self, k):
        try:
            return self[k]
        except KeyError as e:
            raise AttributeError(k) from e


def load_style(spec_path: str | Path, deck_index: int = 0) -> Style:
    """analyze_references.py が出力した deck-spec.json からスタイルを読む。"""
    data = json.loads(Path(spec_path).read_text(encoding="utf-8"))
    deck = data["decks"][deck_index]
    theme = deck.get("theme", {})
    fonts = theme.get("fonts", {})
    major = fonts.get("majorFont", {}) or {}
    minor = fonts.get("minorFont", {}) or {}
    used = [c.split(":", 1)[1] for c, _ in deck.get("stats", {}).get("colors", [])
            if re.fullmatch(r"(?:text|fill):[0-9A-F]{6}", c)]
    return Style(
        file=deck.get("file"),
        canvas=(deck.get("slide_w_in"), deck.get("slide_h_in")),
        theme_colors={k: v for k, v in (theme.get("colors") or {}).items()},
        used_colors=list(dict.fromkeys(used)),
        latin=major.get("latin") or minor.get("latin") or "Century Gothic",
        ea=major.get("ea") or minor.get("ea") or "Meiryo UI",
        layouts=[l["name"] for l in deck.get("layouts", [])],
        sizes=[s for s, _ in deck.get("stats", {}).get("font_sizes", [])],
    )


def rgb(value) -> RGBColor:
    """'1B3928' / '#1B3928' / RGBColor いずれも受け付ける。"""
    if isinstance(value, RGBColor):
        return value
    return RGBColor.from_string(str(value).lstrip("#").upper())


# --------------------------------------------------------------------------- deck
def open_template(path: str | Path, keep_slides: bool = False) -> Presentation:
    """参照 pptx をテンプレートとして開く（既定で既存スライドを全削除）。"""
    prs = Presentation(str(path))
    if not keep_slides:
        sld_id_lst = prs.slides._sldIdLst
        for sld_id in list(sld_id_lst):
            prs.part.drop_rel(sld_id.get(qn("r:id")))
            sld_id_lst.remove(sld_id)
    return prs


def new_from_scratch(width_in: float = 13.333, height_in: float = 7.5) -> Presentation:
    """テンプレートを使わずに空デッキを作る（参照デッキが無い場合の退避用）。"""
    prs = Presentation()
    prs.slide_width = Inches(width_in)
    prs.slide_height = Inches(height_in)
    return prs


def save(prs: Presentation, path: str | Path) -> Path:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


# --------------------------------------------------------------------------- layout
def list_layouts(prs: Presentation) -> list[str]:
    return [l.name for m in prs.slide_masters for l in m.slide_layouts]


def layout_by_name(prs: Presentation, name: str | int):
    """レイアウトを名前（完全一致→部分一致）またはインデックスで取得する。"""
    layouts = [l for m in prs.slide_masters for l in m.slide_layouts]
    if isinstance(name, int):
        return layouts[name]
    for l in layouts:
        if l.name == name:
            return l
    for l in layouts:
        if name.lower() in l.name.lower():
            return l
    raise KeyError("レイアウト '%s' が見つかりません。候補: %s" % (name, [l.name for l in layouts]))


def add_slide(prs: Presentation, layout: str | int = 0):
    return prs.slides.add_slide(layout_by_name(prs, layout))


def set_placeholder(slide, key, text: str, size: float | None = None,
                    bold: bool | None = None, color=None,
                    latin: str | None = None, ea: str | None = None):
    """プレースホルダに文字を入れる。key は idx(int) か型名の部分文字列('TITLE' 等)。

    書式を指定しなければレイアウト／マスターの継承書式がそのまま効く。
    """
    target = None
    for ph in slide.placeholders:
        fmt = ph.placeholder_format
        if isinstance(key, int):
            if fmt.idx == key:
                target = ph
                break
        elif str(key).upper() in str(fmt.type).upper():
            target = ph
            break
    if target is None:
        raise KeyError("プレースホルダ '%s' がありません。存在するのは %s"
                       % (key, [(p.placeholder_format.idx, str(p.placeholder_format.type)) for p in slide.placeholders]))

    tf = target.text_frame
    tf.text = text
    for para in tf.paragraphs:
        for run in para.runs:
            if size is not None:
                run.font.size = Pt(size)
            if bold is not None:
                run.font.bold = bold
            if color is not None:
                run.font.color.rgb = rgb(color)
            if latin or ea:
                set_run_fonts(run, latin, ea)
    return target


def drop_empty_placeholders(slide) -> int:
    """未使用のプレースホルダ（"クリックしてテキストを入力"）を消す。"""
    removed = 0
    for ph in list(slide.placeholders):
        if not ph.has_text_frame or not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)
            removed += 1
    return removed


def set_slide_background(slide, color) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


# --------------------------------------------------------------------------- text
def set_run_fonts(run, latin: str | None = None, ea: str | None = None) -> None:
    """1 run に欧文フォント(a:latin)と日本語フォント(a:ea)を別々に設定する。

    python-pptx の run.font.name は a:latin しか書かないため、日本語が
    欧文フォントにフォールバックして字形が崩れる。本関数で a:ea を明示する。
    """
    rPr = run._r.get_or_add_rPr()
    if latin:
        rPr.get_or_add_latin().set("typeface", latin)
    if ea:
        el = rPr.find(qn("a:ea"))
        if el is None:
            el = rPr.makeelement(qn("a:ea"), {})
            rPr.insert_element_before(el, "a:cs", "a:sym", "a:hlinkClick",
                                      "a:hlinkMouseOver", "a:rtl", "a:extLst")
        el.set("typeface", ea)


def _style_run(run, size, bold, italic, color, latin, ea):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)
    set_run_fonts(run, latin, ea)


def add_text(slide, text, x, y, w, h, *, size=14, bold=None, italic=None, color="404040",
             align="l", anchor="t", latin="Century Gothic", ea="Meiryo UI",
             line_spacing=None, space_after=0, wrap=True, margin=0.0, shrink=False):
    """テキストボックスを置く。text は str か、run 単位の dict のリスト。

    dict 形式の例: [{"text": "重要", "bold": True, "color": "1B3928"}, {"text": "な点"}]
    改行は "\\n" で段落分割される。shrink=True で枠に収まるようサイズを自動縮小。
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = ANCHOR[anchor]
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, Inches(margin))

    runs = [{"text": text}] if isinstance(text, str) else list(text)
    if shrink:
        plain = "".join(r.get("text", "") for r in runs)
        size = fit_size(plain, w - margin * 2, h - margin * 2, size)

    first_para = True
    para = tf.paragraphs[0]
    for spec in runs:
        chunks = str(spec.get("text", "")).split("\n")
        for i, chunk in enumerate(chunks):
            if i > 0 or (not first_para and spec.get("newline")):
                para = tf.add_paragraph()
            first_para = False
            para.alignment = ALIGN[spec.get("align", align)]
            if line_spacing:
                para.line_spacing = line_spacing
            para.space_after = Pt(spec.get("space_after", space_after))
            if chunk == "":
                continue
            run = para.add_run()
            run.text = chunk
            _style_run(run, spec.get("size", size), spec.get("bold", bold),
                       spec.get("italic", italic), spec.get("color", color),
                       spec.get("latin", latin), spec.get("ea", ea))
    return box


def add_bullets(slide, items, x, y, w, h, *, size=14, color="404040", bullet_char="・",
                indent_in=0.22, latin="Century Gothic", ea="Meiryo UI",
                line_spacing=1.3, space_after=4, anchor="t"):
    """箇条書き。items は str か (level, text) タプル、または dict。

    PowerPoint の自動 bullet ではなく文字を前置する方式（日本語資料で崩れにくい）。
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ANCHOR[anchor]
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, Inches(0))

    norm = []
    for it in items:
        if isinstance(it, dict):
            norm.append((it.get("level", 0), it.get("text", ""), it))
        elif isinstance(it, (tuple, list)):
            norm.append((it[0], it[1], {}))
        else:
            norm.append((0, it, {}))

    for i, (level, text, opt) in enumerate(norm):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.line_spacing = line_spacing
        para.space_after = Pt(opt.get("space_after", space_after))
        marker = opt.get("bullet", bullet_char if level == 0 else "－")
        prefix = " " * 0
        run = para.add_run()
        run.text = "%s%s%s" % (prefix, (marker + " ") if marker else "", text)
        _style_run(run, opt.get("size", size - level), opt.get("bold"), None,
                   opt.get("color", color), latin, ea)
        if level:
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(indent_in * level))))
            pPr.set("indent", "0")
    return box


# --------------------------------------------------------------------------- shapes
_SHADOW_TAGS = ("a:outerShdw", "a:innerShdw", "a:prstShdw", "a:reflection", "a:glow", "a:softEdge")


def no_shadow(shape):
    """図形から影・光彩系の効果を構造的に取り除く（**禁じ手の強制**）。

    効果指定は 2 系統あり、両方を潰さないと影は消えない:
      1. `<p:spPr><a:effectLst>` … 図形に直接付いた効果
      2. `<p:style><a:effectRef idx="N">` … テーマ effectStyleLst の N 番を参照

    python-pptx の `shape.shadow.inherit = False` は 1 に空要素を書くだけで
    2 の effectRef は残る。PowerPoint は空の effectLst を優先するが、
    LibreOffice 等は effectRef を拾って影を描くため、ここで effectRef を
    idx="0"（効果なし）に倒し、全レンダラで影が出ない状態にする。
    """
    el = shape._element
    sp_pr = el.find(qn("p:spPr"))
    if sp_pr is not None:
        for lst in sp_pr.findall(qn("a:effectLst")):
            for child in list(lst):
                lst.remove(child)
        for dag in sp_pr.findall(qn("a:effectDag")):
            sp_pr.remove(dag)
    style = el.find(qn("p:style"))
    if style is not None:
        for ref in style.findall(qn("a:effectRef")):
            ref.set("idx", "0")
            for child in list(ref):
                ref.remove(child)
    return shape


def add_rect(slide, x, y, w, h, *, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE,
             text=None, size=12, bold=None, color="404040", align="c", anchor="m",
             latin="Century Gothic", ea="Meiryo UI", margin=0.05, shadow=False):
    """矩形（角丸・矢印等も MSO_SHAPE で指定可）。text を渡せば中に文字も入れる。

    影は禁じ手なので付けられない（`shadow=True` は ValueError）。
    """
    if shadow:
        raise ValueError(
            "ドロップシャドウは本スキルの禁止パターンです（SKILL.md 禁止パターン）。"
            "階層は矩形＋境界線（line=）か背景色（fill=）で表現してください。")
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    no_shadow(sp)  # テーマ effectRef 由来の影まで落とす

    tf = sp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ANCHOR[anchor]
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, Inches(margin))
    if text is None:
        return sp
    runs = [{"text": text}] if isinstance(text, str) else list(text)
    first = True
    para = tf.paragraphs[0]
    for spec in runs:
        for i, chunk in enumerate(str(spec.get("text", "")).split("\n")):
            if i > 0 or not first:
                para = tf.add_paragraph()
            first = False
            para.alignment = ALIGN[spec.get("align", align)]
            if chunk == "":
                continue
            run = para.add_run()
            run.text = chunk
            _style_run(run, spec.get("size", size), spec.get("bold", bold), None,
                       spec.get("color", color), spec.get("latin", latin), spec.get("ea", ea))
    return sp


def add_line(slide, x1, y1, x2, y2, *, color="D0D0D0", width=0.75, dash=None):
    from pptx.enum.shapes import MSO_CONNECTOR

    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                      Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width)
    if dash:
        conn.line.dash_style = dash
    return no_shadow(conn)


def add_table(slide, data, x, y, w, h, *, col_widths=None, row_heights=None,
              header=True, header_fill="1B3928", header_color="FFFFFF",
              body_color="404040", border_color="D9D9D9", band_fill=None,
              size=11, header_size=None, latin="Century Gothic", ea="Meiryo UI",
              align="l", plain_style=True):
    """2 次元リストから表を作る。data[0] をヘッダ行として扱う（header=False で無効）。"""
    rows, cols = len(data), max(len(r) for r in data)
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    if plain_style:
        tbl_pr = tbl._tbl.get_or_add_tblPr()
        tbl_pr.set("firstRow", "1" if header else "0")
        tbl_pr.set("bandRow", "1" if band_fill else "0")
        style_el = tbl_pr.find(qn("a:tableStyleId"))
        if style_el is None:
            style_el = tbl_pr.makeelement(qn("a:tableStyleId"), {})
            tbl_pr.append(style_el)
        style_el.text = _PLAIN_TABLE_STYLE

    if col_widths:
        for i, cw in enumerate(col_widths[:cols]):
            tbl.columns[i].width = Inches(cw)
    if row_heights:
        for i, rh in enumerate(row_heights[:rows]):
            tbl.rows[i].height = Inches(rh)

    for r, row in enumerate(data):
        is_head = header and r == 0
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = str(row[c]) if c < len(row) else ""
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if is_head:
                cell.fill.fore_color.rgb = rgb(header_fill)
            elif band_fill and r % 2 == 0:
                cell.fill.fore_color.rgb = rgb(band_fill)
            else:
                cell.fill.fore_color.rgb = rgb("FFFFFF")
            for para in cell.text_frame.paragraphs:
                para.alignment = ALIGN["c"] if is_head else ALIGN[align]
                for run in para.runs:
                    _style_run(run, header_size or (size if not is_head else size),
                               True if is_head else None, None,
                               header_color if is_head else body_color, latin, ea)
            _cell_borders(cell, border_color)
    return shape


def _cell_borders(cell, color) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        el = tc_pr.find(qn(tag))
        if el is None:
            el = tc_pr.makeelement(qn(tag), {"w": "6350", "cap": "flat", "cmpd": "sng", "algn": "ctr"})
            tc_pr.insert(0, el)
        for child in list(el):
            el.remove(child)
        fill = el.makeelement(qn("a:solidFill"), {})
        clr = el.makeelement(qn("a:srgbClr"), {"val": str(color).lstrip("#").upper()})
        fill.append(clr)
        el.append(fill)


def add_picture(slide, image_path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return no_shadow(slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), **kw))


# --------------------------------------------------------------------------- fitting
def text_width_em(text: str) -> float:
    """文字列の概算幅を em 単位で返す（全角=1.0em、半角=0.55em）。"""
    total = 0.0
    for ch in text:
        if ch == "\n":
            continue
        total += 1.0 if ord(ch) > 0x2E7F else 0.55
    return total


def fit_size(text: str, w_in: float, h_in: float, base_pt: float,
             min_pt: float = 8.0, line_spacing: float = 1.25) -> float:
    """枠 (w_in × h_in) に収まる最大のフォントサイズを概算で返す。"""
    size = base_pt
    while size > min_pt:
        chars_per_line = max((w_in * 72.0) / size, 1.0)
        lines = 0
        for para in str(text).split("\n"):
            lines += max(1, -(-text_width_em(para) // chars_per_line))
        if lines * size * line_spacing / 72.0 <= h_in:
            return round(size, 1)
        size -= 0.5
    return min_pt
