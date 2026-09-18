#!/usr/bin/env python3
"""生成した pptx を機械チェックする（はみ出し・溢れ・フォント逸脱・配色逸脱）。

使い方:
    python3 verify_deck.py output.pptx [--spec deck-spec.json] [--min-pt 9] [--strict]

--spec を渡すと、参照デッキから抽出したフォント／配色と突き合わせて逸脱を報告する。
終了コード: 0=ERROR なし / 1=ERROR あり（--strict では WARN でも 1）。
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from deck_kit import fit_size, text_width_em  # noqa: E402


def inch(v):
    return round(Emu(int(v)).inches, 3) if v is not None else None


def iter_shapes(shapes, depth=0):
    for sh in shapes:
        yield sh
        if depth < 3 and sh.shape_type is not None and "GROUP" in str(sh.shape_type):
            yield from iter_shapes(sh.shapes, depth + 1)


SHADOW_TAGS = ("a:outerShdw", "a:innerShdw", "a:prstShdw", "a:reflection", "a:glow")


def shadow_of(shape):
    """図形に影（等の効果）が付いているかを返す。

    戻り値は None / "直接指定" / "テーマ参照"。
    `<a:effectLst>` 配下の影要素は直接指定。空でも `<p:style><a:effectRef>` の
    idx が 0 以外ならテーマの効果を参照しており、レンダラ次第で影が出る。
    """
    el = getattr(shape, "_element", None)
    if el is None:
        return None
    sp_pr = el.find(qn("p:spPr"))
    if sp_pr is not None:
        for lst in sp_pr.findall(qn("a:effectLst")):
            for tag in SHADOW_TAGS:
                if lst.find(qn(tag)) is not None:
                    return "直接指定"
        if sp_pr.find(qn("a:effectDag")) is not None:
            return "直接指定"
    style = el.find(qn("p:style"))
    if style is not None:
        for ref in style.findall(qn("a:effectRef")):
            if (ref.get("idx") or "0") != "0":
                return "テーマ参照"
    return None


def effective_runs(shape):
    if not getattr(shape, "has_text_frame", False):
        return []
    out = []
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.text.strip():
                out.append(run)
    return out


def check(pptx_path: Path, spec_path: Path | None, min_pt: float) -> list[tuple]:
    prs = Presentation(str(pptx_path))
    W, H = inch(prs.slide_width), inch(prs.slide_height)

    ref_fonts, ref_colors = set(), set()
    if spec_path and spec_path.is_file():
        data = json.loads(spec_path.read_text(encoding="utf-8"))
        for deck in data.get("decks", []):
            for kind in (deck.get("theme", {}).get("fonts") or {}).values():
                ref_fonts.update(v for v in (kind or {}).values() if v)
            for name, _ in deck.get("stats", {}).get("fonts", []):
                ref_fonts.add(name.split("(")[0])
            ref_colors.update((deck.get("theme", {}).get("colors") or {}).values())
            for c, _ in deck.get("stats", {}).get("colors", []):
                if ":" in c:
                    ref_colors.add(c.split(":", 1)[1])
        ref_fonts = {f for f in ref_fonts if f and not f.startswith("+")}
        ref_colors = {c for c in ref_colors if c and len(c) == 6}

    issues: list[tuple] = []

    def add(level, no, msg):
        issues.append((level, no, msg))

    for i, slide in enumerate(prs.slides, start=1):
        boxes = []
        for sh in iter_shapes(slide.shapes):
            x, y, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
            if None in (x, y, w, h):
                continue
            name = sh.name

            # 1. キャンバスからのはみ出し
            if x < -0.02 or y < -0.02 or x + w > W + 0.02 or y + h > H + 0.02:
                add("ERROR", i, "スライド外にはみ出し: %s (x=%s y=%s w=%s h=%s / canvas %sx%s)"
                    % (name, x, y, w, h, W, H))

            # 1b. 影（禁じ手）: 直接指定は ERROR、テーマ参照は描画依存で出るため WARN
            src = shadow_of(sh)
            if src == "直接指定":
                add("ERROR", i, "影・効果が付いている（禁じ手）: %s — deck_kit.no_shadow() を通すこと" % name)
            elif src == "テーマ参照":
                add("WARN", i, "テーマ効果 (effectRef) が残っている: %s"
                              " — PowerPoint では出ないが LibreOffice 等で影が描かれる" % name)

            runs = effective_runs(sh)
            text = "".join(r.text for r in runs)

            # 2. プレースホルダの空／既定文言残り
            if sh.is_placeholder and not text.strip():
                add("WARN", i, "空のプレースホルダが残っている: %s（drop_empty_placeholders で削除）" % name)
            if "クリックしてテキストを入力" in text or "Click to edit" in text:
                add("ERROR", i, "プレースホルダ既定文言が残っている: %s" % name)

            if runs:
                sizes = [r.font.size.pt for r in runs if r.font.size is not None]
                # 3. 小さすぎる文字
                for s in sizes:
                    if s < min_pt:
                        add("WARN", i, "フォントが小さい (%.1fpt < %.1fpt): %s" % (s, min_pt, name))
                        break
                # 4. テキスト溢れ（概算）
                base = max(sizes) if sizes else None
                if base and w > 0.2 and h > 0.15 and not getattr(sh, "has_table", False):
                    fitted = fit_size(text, w, h, base)
                    if fitted < base - 0.6:
                        add("WARN", i, "枠に収まらない可能性 (%.1fpt 指定 / 推定上限 %.1fpt, %d 文字): %s"
                            % (base, fitted, len(text), name))
                # 5. フォント逸脱
                if ref_fonts:
                    for r in runs:
                        if r.font.name and r.font.name not in ref_fonts:
                            add("WARN", i, "参照デッキに無いフォント '%s': %s" % (r.font.name, name))
                            break
                # 6. 配色逸脱
                if ref_colors:
                    for r in runs:
                        try:
                            c = str(r.font.color.rgb).upper() if r.font.color.type is not None else None
                        except Exception:
                            c = None
                        if c and c not in ref_colors and c not in ("FFFFFF", "000000"):
                            add("INFO", i, "参照デッキに無いテキスト色 #%s: %s" % (c, name))
                            break
                boxes.append((x, y, w, h, name))

        # 7. テキストボックス同士の重なり
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                ax, ay, aw, ah, an = boxes[a]
                bx, by, bw, bh, bn = boxes[b]
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 0.08 and oy > 0.08:
                    add("INFO", i, "テキスト枠が重なっている: %s × %s (%.2f×%.2f inch)" % (an, bn, ox, oy))

    if not prs.slides:
        add("ERROR", 0, "スライドが 1 枚もありません")
    return issues


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--spec", default=None, help="analyze_references.py が出力した deck-spec.json")
    ap.add_argument("--min-pt", type=float, default=9.0)
    ap.add_argument("--strict", action="store_true", help="WARN も失敗扱いにする")
    args = ap.parse_args()

    issues = check(Path(args.pptx), Path(args.spec) if args.spec else None, args.min_pt)
    counts = Counter(lv for lv, _, _ in issues)
    order = {"ERROR": 0, "WARN": 1, "INFO": 2}
    for lv, no, msg in sorted(issues, key=lambda t: (order[t[0]], t[1])):
        print("[%-5s] slide %-3s %s" % (lv, no, msg))
    print("\n-- ERROR %d / WARN %d / INFO %d --"
          % (counts["ERROR"], counts["WARN"], counts["INFO"]))
    if counts["ERROR"] or (args.strict and counts["WARN"]):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
